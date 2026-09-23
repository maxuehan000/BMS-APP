# -*- coding: utf-8 -*-
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = r"F:\_proto_tmp.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
for ws in wb.worksheets:
    print("=" * 80)
    print("SHEET:", ws.title, "dims:", ws.dimensions, "max_row:", ws.max_row, "max_col:", ws.max_column)
    print("=" * 80)
    for row in ws.iter_rows():
        cells = []
        for c in row:
            v = c.value
            if v is None:
                continue
            s = str(v).replace("\n", "\\n")
            cells.append(f"{c.coordinate}={s}")
        if cells:
            print(" | ".join(cells))
